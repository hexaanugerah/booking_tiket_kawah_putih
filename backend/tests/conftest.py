"""Test configuration.

Environment variables are set before the application is imported so tests use a
throwaway database and storage directory instead of development data.
"""
from __future__ import annotations

import io
import os
from pathlib import Path
import tempfile
import uuid

import pytest

_TEST_ROOT = Path(tempfile.mkdtemp(prefix="docuconvert-tests-"))
os.environ.update(
    {
        "ENVIRONMENT": "test",
        "DATABASE_URL": f"sqlite:///{(_TEST_ROOT / 'test.db').as_posix()}",
        "STORAGE_DIR": str(_TEST_ROOT / "storage"),
        "SECRET_KEY": "test-secret-key-not-for-production-use",
        "RATE_LIMIT_PER_MINUTE": "10000",
        "ALLOW_ANONYMOUS": "true",
        "RESULT_TTL_HOURS": "1",
        "CLEANUP_INTERVAL_MINUTES": "60",
    }
)


@pytest.fixture()
def client():
    from fastapi.testclient import TestClient

    from app.main import app

    with TestClient(app) as test_client:
        yield test_client


@pytest.fixture()
def pdf_bytes():
    return make_pdf()


@pytest.fixture()
def docx_bytes():
    return make_docx()


@pytest.fixture()
def xlsx_bytes():
    return make_xlsx()


@pytest.fixture()
def png_bytes():
    return make_png()


def make_pdf(text: str = "Quarterly Report", pages: int = 1) -> bytes:
    import fitz

    document = fitz.open()
    for page_number in range(pages):
        page = document.new_page(width=595, height=842)
        page.insert_text((72, 70), text, fontsize=22, fontname="helv")
        page.insert_text(
            (72, 110),
            f"Page {page_number + 1} body text with details about the report.",
            fontsize=11,
        )
        page.insert_text((72, 132), "1. Revenue grew by 12 percent.", fontsize=11)
        page.insert_text((72, 154), "2. Costs stayed flat.", fontsize=11)
        if page_number == 0:
            page.draw_rect(fitz.Rect(72, 180, 500, 240), color=(0, 0, 0))
            page.draw_line(fitz.Point(300, 180), fitz.Point(300, 240), color=(0, 0, 0))
            for y in (210, 230):
                page.draw_line(fitz.Point(72, y), fitz.Point(500, y), color=(0, 0, 0))
            page.insert_text((80, 196), "Item", fontsize=10)
            page.insert_text((310, 196), "Amount", fontsize=10)
            page.insert_text((80, 226), "Widgets", fontsize=10)
            page.insert_text((310, 226), "1200", fontsize=10)
    data = document.tobytes()
    document.close()
    return data


def make_docx(text: str = "Sample Word document") -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("Test Heading", 1)
    document.add_paragraph(text)
    document.add_paragraph("First bullet", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Score"
    table.cell(1, 0).text = "Ani"
    table.cell(1, 1).text = "95"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_xlsx() -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Name", "Score"])
    sheet.append(["Ani", 95])
    sheet.append(["Budi", 88])
    second = workbook.create_sheet("Summary")
    second.append(["Total", 183])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_png(text: str = "HELLO DOCUCONVERT") -> bytes:
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (1100, 400), "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("arial.ttf", 80)
    except Exception:
        font = ImageFont.load_default()
    draw.text((40, 140), text, fill="black", font=font)
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def upload(client, tool: str, files: list[tuple[str, bytes]], options: dict | None = None):
    import json

    return client.post(
        "/api/v1/conversions",
        data={"tool": tool, "options": json.dumps(options or {})},
        files=[("files", (name, content, "application/octet-stream")) for name, content in files],
    )


def unique_email() -> str:
    return f"user-{uuid.uuid4().hex[:10]}@example.com"
