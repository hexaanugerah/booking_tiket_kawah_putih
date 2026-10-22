"""PDF -> XLSX table extraction using pdfplumber and openpyxl."""
from __future__ import annotations

from pathlib import Path
import re

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.errors import ConversionError
from app.services.base import ConversionContext, ConversionResult
from app.services.pdf_ops import _open_pdf

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_PAGES = 150
_INT_RE = re.compile(r"^-?\d{1,15}$")
_FLOAT_RE = re.compile(r"^-?\d{1,15}\.\d+$")


def _to_number(value: str):
    text = value.strip()
    if not text:
        return value
    if _INT_RE.match(text):
        if len(text.lstrip("-")) > 1 and text.lstrip("-").startswith("0"):
            return value
        return int(text)
    if _FLOAT_RE.match(text):
        return float(text)
    return value


def _clean_sheet_name(name: str, used: set[str]) -> str:
    name = re.sub(r"[\[\]:*?/\\]", "-", name).strip()[:31] or "Sheet"
    candidate = name
    index = 2
    while candidate in used:
        suffix = f" ({index})"
        candidate = f"{name[:31 - len(suffix)]}{suffix}"
        index += 1
    used.add(candidate)
    return candidate


def _write_table(wb: Workbook, title: str, rows: list[list[str]], used: set[str]) -> None:
    sheet = wb.create_sheet(_clean_sheet_name(title, used))
    widths: dict[int, int] = {}
    for row_index, row in enumerate(rows, start=1):
        for column_index, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=_to_number(value))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if row_index == 1:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="E8F0FE")
            widths[column_index] = max(widths.get(column_index, 8), min(60, len(str(value)) + 2))
    for column_index, width in widths.items():
        sheet.column_dimensions[get_column_letter(column_index)].width = width
    sheet.freeze_panes = "A2"


def convert(ctx: ConversionContext) -> ConversionResult:
    _open_pdf(ctx.input).close()  # validates the file and encryption status
    text_fallback = bool(ctx.options.get("text_fallback", True))
    warnings: list[str] = []
    wb = Workbook()
    wb.remove(wb.active)
    used_names: set[str] = set()
    tables_found = 0
    text_rows: list[list[str]] = []

    try:
        with pdfplumber.open(str(ctx.input)) as pdf:
            page_count = len(pdf.pages)
            if page_count > MAX_PAGES:
                warnings.append(f"Only the first {MAX_PAGES} pages were analysed.")
            for page_index, page in enumerate(pdf.pages[:MAX_PAGES]):
                try:
                    tables = page.extract_tables()
                except Exception:
                    tables = []
                for table_index, rows in enumerate(tables, start=1):
                    cleaned = []
                    for row in rows:
                        values = ["" if cell is None else str(cell).strip() for cell in row]
                        if any(values):
                            cleaned.append(values)
                    width = max((len(row) for row in cleaned), default=0)
                    cleaned = [row + [""] * (width - len(row)) for row in cleaned]
                    if len(cleaned) < 2 or width < 2:
                        continue
                    tables_found += 1
                    _write_table(wb, f"Page {page_index + 1} T{table_index}", cleaned, used_names)
                if text_fallback:
                    for line in (page.extract_text() or "").splitlines():
                        line = line.strip()
                        if line:
                            text_rows.append([line])
    except pdfplumber.PDFSyntaxError as exc:
        raise ConversionError("The PDF could not be parsed. It may be damaged.") from exc

    if tables_found == 0:
        if not text_fallback:
            raise ConversionError(
                "No tables were detected in this PDF. Try a PDF with ruled tables or enable the text fallback."
            )
        if not text_rows:
            raise ConversionError("No extractable table or text content was found in this PDF.")
        warnings.append("No ruled tables were detected; plain text lines were exported instead.")
        _write_table(wb, "Text", text_rows, used_names)
    elif text_fallback and text_rows:
        warnings.append("Only detected tables are included; remaining text lines were not exported.")

    if not wb.sheetnames:
        raise ConversionError("No exportable content was found in this PDF.")

    output = ctx.workdir / "tables.xlsx"
    wb.save(output)
    return ConversionResult(output, "tables.xlsx", XLSX_MIME, warnings)
