"""End-to-end conversion tests through the HTTP API."""
from __future__ import annotations

import io
import zipfile

import pytest

from tests.conftest import make_docx, make_pdf, make_png, make_xlsx, upload


def _completed(client, response):
    assert response.status_code == 202, response.text
    job_id = response.json()["id"]
    detail = client.get(f"/api/v1/conversions/{job_id}")
    assert detail.status_code == 200, detail.text
    data = detail.json()
    assert data["status"] == "completed", data.get("error_message")
    return data


def _download(client, job):
    assert job["download_url"]
    response = client.get(job["download_url"])
    assert response.status_code == 200, response.text
    return response.content


def test_pdf_to_word_produces_editable_document(client):
    from docx import Document

    job = _completed(client, upload(client, "pdf-to-word", [("report.pdf", make_pdf())]))
    content = _download(client, job)
    document = Document(io.BytesIO(content))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Quarterly Report" in text
    assert "Revenue grew" in text
    assert len(document.tables) == 1
    assert document.tables[0].cell(0, 0).text.strip() == "Item"
    # The result must be editable text, not a single full-page image.
    assert len(document.inline_shapes) == 0
    assert document.paragraphs[0].style.name.startswith("Heading")
    assert job["warnings"]


def test_pdf_to_excel_extracts_tables(client):
    from openpyxl import load_workbook

    job = _completed(client, upload(client, "pdf-to-excel", [("report.pdf", make_pdf())]))
    workbook = load_workbook(io.BytesIO(_download(client, job)))
    assert workbook.sheetnames
    values = [
        [cell for cell in row]
        for row in workbook[workbook.sheetnames[0]].iter_rows(values_only=True)
    ]
    assert values[0] == ["Item", "Amount"]
    assert values[1][0] == "Widgets"
    assert values[1][1] == 1200


def test_pdf_to_pptx(client):
    from pptx import Presentation

    job = _completed(client, upload(client, "pdf-to-pptx", [("report.pdf", make_pdf())]))
    presentation = Presentation(io.BytesIO(_download(client, job)))
    assert len(presentation.slides) == 1
    slide_text = "\n".join(
        shape.text_frame.text for shape in presentation.slides[0].shapes if shape.has_text_frame
    )
    assert "Quarterly Report" in slide_text


def test_pdf_to_images_returns_zip(client):
    job = _completed(
        client,
        upload(client, "pdf-to-images", [("report.pdf", make_pdf(pages=2))], options={"format": "png", "dpi": 100}),
    )
    with zipfile.ZipFile(io.BytesIO(_download(client, job))) as archive:
        names = archive.namelist()
        assert len(names) == 2
        assert archive.read(names[0]).startswith(b"\x89PNG")


def test_word_to_pdf(client):
    job = _completed(client, upload(client, "docx-to-pdf", [("sample.docx", make_docx())]))
    content = _download(client, job)
    assert content.startswith(b"%PDF")
    assert len(content) > 500


def test_word_to_txt_html_and_markdown(client):
    txt_job = _completed(client, upload(client, "docx-to-txt", [("sample.docx", make_docx())]))
    text = _download(client, txt_job).decode("utf-8")
    assert "Test Heading" in text
    assert "First bullet" in text

    html_job = _completed(client, upload(client, "docx-to-html", [("sample.docx", make_docx())]))
    html = _download(client, html_job).decode("utf-8")
    assert "<h1" in html and "Test Heading" in html

    md_job = _completed(client, upload(client, "docx-to-markdown", [("sample.docx", make_docx())]))
    markdown = _download(client, md_job).decode("utf-8")
    assert "# " in markdown
    assert "| Name | Score |" in markdown


def test_txt_to_word_and_markdown_to_word(client):
    job = _completed(
        client,
        upload(client, "txt-to-docx", [("notes.txt", b"First paragraph.\n\nSecond paragraph.")]),
    )
    from docx import Document

    document = Document(io.BytesIO(_download(client, job)))
    assert "First paragraph." in document.paragraphs[0].text

    markdown = b"# Title\n\nSome **bold** text.\n\n- one\n- two\n\n| A | B |\n| --- | --- |\n| 1 | 2 |\n"
    md_job = _completed(client, upload(client, "markdown-to-docx", [("doc.md", markdown)]))
    document = Document(io.BytesIO(_download(client, md_job)))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Title" in text
    assert "bold" in text
    assert len(document.tables) == 1


def test_word_to_word_merge(client):
    job = _completed(
        client,
        upload(
            client,
            "merge-docx",
            [("one.docx", make_docx("First document")), ("two.docx", make_docx("Second document"))],
        ),
    )
    from docx import Document

    document = Document(io.BytesIO(_download(client, job)))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert text.count("Test Heading") == 2
    assert "Second document" in text


def test_clean_word(client):
    job = _completed(client, upload(client, "clean-docx", [("sample.docx", make_docx())]))
    from docx import Document

    document = Document(io.BytesIO(_download(client, job)))
    assert document.paragraphs


def test_excel_to_csv_multiple_sheets(client):
    job = _completed(client, upload(client, "xlsx-to-csv", [("data.xlsx", make_xlsx())]))
    content = _download(client, job)
    assert job["output_filename"].endswith(".zip")
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        names = archive.namelist()
        assert "Data.csv" in names and "Summary.csv" in names
        assert "Ani,95" in archive.read("Data.csv").decode("utf-8-sig")


def test_excel_to_csv_single_sheet(client):
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.append(["A", "B"])
    workbook.active.append([1, 2])
    buffer = io.BytesIO()
    workbook.save(buffer)

    job = _completed(client, upload(client, "xlsx-to-csv", [("single.xlsx", buffer.getvalue())]))
    assert job["output_filename"] == "single.csv"
    assert "A,B" in _download(client, job).decode("utf-8-sig")


def test_csv_to_excel(client):
    from openpyxl import load_workbook

    job = _completed(
        client,
        upload(client, "csv-to-xlsx", [("scores.csv", b"Name,Score\nAni,95\nBudi,88\n")]),
    )
    workbook = load_workbook(io.BytesIO(_download(client, job)))
    sheet = workbook[workbook.sheetnames[0]]
    assert sheet["A1"].value == "Name"
    assert sheet["B2"].value == 95


def test_excel_to_html_and_pdf(client):
    html_job = _completed(client, upload(client, "xlsx-to-html", [("data.xlsx", make_xlsx())]))
    html = _download(client, html_job).decode("utf-8")
    assert "<table>" in html and "Budi" in html

    pdf_job = _completed(client, upload(client, "xlsx-to-pdf", [("data.xlsx", make_xlsx())]))
    assert _download(client, pdf_job).startswith(b"%PDF")


def test_image_to_pdf_and_format_conversion(client):
    pdf_job = _completed(client, upload(client, "image-to-pdf", [("scan.png", make_png())]))
    assert _download(client, pdf_job).startswith(b"%PDF")

    jpg_job = _completed(
        client,
        upload(client, "convert-image", [("scan.png", make_png())], options={"target": "jpg", "quality": 80}),
    )
    assert jpg_job["output_filename"].endswith(".jpg")
    assert _download(client, jpg_job).startswith(b"\xff\xd8\xff")


def _require_ocr():
    from app.services.capabilities import ocr_available

    if not ocr_available():
        pytest.skip("No OCR engine is installed on this machine.")


def test_ocr_image(client):
    _require_ocr()
    job = _completed(client, upload(client, "image-ocr", [("scan.png", make_png())]))
    text = _download(client, job).decode("utf-8")
    assert "HELLO" in text.upper()


def test_pdf_merge_and_split(client):
    first = make_pdf("First Document", pages=1)
    second = make_pdf("Second Document", pages=2)
    merged = _completed(client, upload(client, "merge-pdf", [("a.pdf", first), ("b.pdf", second)]))
    merged_bytes = _download(client, merged)

    import fitz

    document = fitz.open(stream=merged_bytes, filetype="pdf")
    assert document.page_count == 3
    document.close()

    split = _completed(
        client, upload(client, "split-pdf", [("merged.pdf", merged_bytes)], options={"mode": "each"})
    )
    with zipfile.ZipFile(io.BytesIO(_download(client, split))) as archive:
        assert len(archive.namelist()) == 3


def test_pdf_rotate_page_numbers_watermark_compress_remove_reorder(client):
    pdf = make_pdf(pages=2)

    rotated = _completed(
        client, upload(client, "rotate-pdf", [("a.pdf", pdf)], options={"angle": "90", "pages": "1"})
    )
    import fitz

    document = fitz.open(stream=_download(client, rotated), filetype="pdf")
    assert document[0].rotation == 90
    assert document[1].rotation == 0
    document.close()

    numbered = _completed(client, upload(client, "add-page-numbers", [("a.pdf", pdf)]))
    document = fitz.open(stream=_download(client, numbered), filetype="pdf")
    assert "1" in document[0].get_text()
    document.close()

    watermarked = _completed(
        client,
        upload(client, "add-watermark", [("a.pdf", pdf)], options={"text": "INTERNAL", "opacity": 20}),
    )
    document = fitz.open(stream=_download(client, watermarked), filetype="pdf")
    assert "INTERNAL" in document[0].get_text()
    document.close()

    compressed = _completed(client, upload(client, "compress-pdf", [("a.pdf", pdf)]))
    assert _download(client, compressed).startswith(b"%PDF")

    removed = _completed(
        client, upload(client, "remove-pages", [("a.pdf", pdf)], options={"pages": "1"})
    )
    document = fitz.open(stream=_download(client, removed), filetype="pdf")
    assert document.page_count == 1
    document.close()

    reordered = _completed(
        client, upload(client, "reorder-pdf", [("a.pdf", pdf)], options={"order": "2,1"})
    )
    document = fitz.open(stream=_download(client, reordered), filetype="pdf")
    assert document.page_count == 2
    document.close()


def test_pdf_text_and_table_extraction(client):
    text_job = _completed(
        client,
        upload(client, "extract-text", [("report.pdf", make_pdf())], options={"language": "eng"}),
    )
    text = _download(client, text_job).decode("utf-8")
    assert "Quarterly Report" in text
    assert "Page 1" in text

    table_job = _completed(client, upload(client, "extract-tables", [("report.pdf", make_pdf())]))
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(_download(client, table_job)))
    assert workbook.sheetnames


def test_html_and_markdown_to_pdf(client):
    pdf_job = _completed(
        client,
        upload(client, "html-to-pdf", [("page.html", b"<html><body><h1>Hello</h1><p>World</p></body></html>")]),
    )
    assert _download(client, pdf_job).startswith(b"%PDF")

    md_job = _completed(
        client,
        upload(client, "markdown-to-pdf", [("doc.md", b"# Title\n\nA paragraph.\n\n- item")]),
    )
    assert _download(client, md_job).startswith(b"%PDF")


def test_pdf_ocr_to_word(client):
    _require_ocr()
    import fitz
    from PIL import Image

    image = Image.new("RGB", (1000, 500), "white")
    from PIL import ImageDraw

    draw = ImageDraw.Draw(image)
    draw.text((60, 200), "SCANNED INVOICE 12345", fill="black")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    document = fitz.open()
    page = document.new_page(width=595, height=842)
    page.insert_image(fitz.Rect(40, 40, 555, 500), stream=buffer.getvalue())
    scanned = document.tobytes()
    document.close()

    job = _completed(client, upload(client, "pdf-ocr-to-word", [("scanned.pdf", scanned)]))
    from docx import Document

    output = Document(io.BytesIO(_download(client, job)))
    text = " ".join(paragraph.text for paragraph in output.paragraphs).upper()
    assert "SCANNED" in text
