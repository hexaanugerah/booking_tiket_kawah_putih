"""PDF -> DOCX with real text extraction, heading/list/table detection and OCR fallback.

The output is an editable Word document built from parsed content, never a set of
full-page screenshots. When layout fidelity cannot be guaranteed the converter
reports warnings alongside the result.
"""
from __future__ import annotations

from collections import Counter
import io
from pathlib import Path
import re
from statistics import median

import fitz
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches, Pt

from app.core.errors import ConversionError
from app.services.base import ConversionContext, ConversionResult
from app.services.capabilities import ocr_available
from app.services.pdf_ops import _open_pdf
from app.services.ocr import ocr_pdf_page

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MAX_PAGES = 200

BULLET_RE = re.compile(r"^\s*([\u2022\u25cf\u25aa\u25e6\u2023\u2043\u00b7\-–—*])\s+(.+)$")
NUMBER_RE = re.compile(r"^\s*((?:\d{1,3}|[a-zA-Z])[.)])\s+(.+)$")
HEADING_STYLE = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3"}


def _is_bold(span: dict) -> bool:
    return bool(span.get("flags", 0) & 16) or "bold" in span.get("font", "").lower()


def _is_italic(span: dict) -> bool:
    return bool(span.get("flags", 0) & 2) or "italic" in span.get("font", "").lower()


def _is_superscript(span: dict) -> bool:
    return bool(span.get("flags", 0) & 1)


def _clean(text: str) -> str:
    return re.sub(r"[ \t\u00a0]+", " ", text).strip()


def _line_text(line: dict) -> str:
    parts: list[str] = []
    previous: dict | None = None
    for span in line.get("spans", []):
        text = span.get("text", "")
        if not text:
            continue
        if previous is not None and parts:
            gap = span["bbox"][0] - previous["bbox"][2]
            if gap > max(1.2, span.get("size", 10) * 0.18) and not parts[-1].endswith(" "):
                parts.append(" ")
        parts.append(text)
        previous = span
    return _clean("".join(parts))


def _line_runs(line: dict) -> list[tuple[str, bool, bool, float, bool]]:
    runs: list[tuple[str, bool, bool, float, bool]] = []
    for span in line.get("spans", []):
        text = span.get("text", "")
        if not text:
            continue
        runs.append(
            (
                text,
                _is_bold(span),
                _is_italic(span),
                float(span.get("size", 11)),
                _is_superscript(span),
            )
        )
    return runs


def _join_line_texts(lines: list[str]) -> str:
    result = ""
    for line in lines:
        line = line.strip()
        if not line:
            continue
        if not result:
            result = line
        elif result.endswith("-") and not result.endswith("--"):
            result = result[:-1] + line
        else:
            result = f"{result} {line}"
    return _clean(result)


def _body_font_size(doc: fitz.Document) -> float:
    sizes: Counter[float] = Counter()
    for page in doc:
        for block in page.get_text("dict").get("blocks", []):
            if block.get("type") != 0:
                continue
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    text = span.get("text", "").strip()
                    if len(text) >= 12:
                        sizes[round(float(span.get("size", 11)) * 2) / 2] += len(text)
    if not sizes:
        return 11.0
    return sizes.most_common(1)[0][0]


def _header_footer_texts(doc: fitz.Document) -> set[str]:
    """Find recurring running headers/footers so they are not duplicated in the body."""
    occurrences: Counter[str] = Counter()
    sampled = 0
    for page in doc:
        height = page.rect.height
        candidates: set[str] = set()
        for block in page.get_text("dict").get("blocks", []):
            if block.get("type") != 0:
                continue
            rect = fitz.Rect(block["bbox"])
            text = _clean(" ".join(_line_text(line) for line in block.get("lines", [])))
            if not text or len(text) > 120:
                continue
            if rect.y1 < height * 0.075 or rect.y0 > height * 0.925:
                candidates.add(text)
        for text in candidates:
            occurrences[text] += 1
        sampled += 1
    if sampled < 3:
        return set()
    threshold = max(3, int(sampled * 0.6))
    return {text for text, count in occurrences.items() if count >= threshold}


def _page_tables(path: Path, page_index: int) -> list[tuple[tuple[float, float, float, float], list[list[str]]]]:
    import pdfplumber

    tables: list[tuple[tuple[float, float, float, float], list[list[str]]]] = []
    with pdfplumber.open(str(path)) as pdf:
        if page_index >= len(pdf.pages):
            return []
        page = pdf.pages[page_index]
        for table in page.find_tables():
            rows = table.extract()
            cleaned_rows = []
            for row in rows:
                values = ["" if cell is None else str(cell).strip() for cell in row]
                if any(value for value in values):
                    cleaned_rows.append(values)
            if len(cleaned_rows) >= 2 and max(len(row) for row in cleaned_rows) >= 2:
                tables.append((tuple(table.bbox), cleaned_rows))
    return tables


def _block_text(block: dict) -> tuple[str, list[tuple[str, bool, bool, float, bool]]]:
    lines = []
    runs: list[tuple[str, bool, bool, float, bool]] = []
    for line in block.get("lines", []):
        text = _line_text(line)
        if text:
            lines.append(text)
            runs.extend(_line_runs(line))
    return _join_line_texts(lines), runs


def _classify(text: str, max_size: float, body_size: float, bold: bool, line_count: int) -> tuple[str, int | None]:
    """Return ('p' | 'h' | 'li', level) for a block of text."""
    bullet = BULLET_RE.match(text)
    numbered = NUMBER_RE.match(text)
    if bullet or numbered:
        return "li", 1 if numbered else 0
    if line_count <= 3 and len(text) <= 120:
        ratio = max_size / body_size if body_size else 1
        if ratio >= 1.45:
            return "h", 1
        if ratio >= 1.24:
            return "h", 2
        if ratio >= 1.1 and bold:
            return "h", 3
        if ratio >= 1.6:
            return "h", 1
    return "p", None


def _add_runs(paragraph, runs: list[tuple[str, bool, bool, float, bool]], *, bold_all: bool = False) -> None:
    for text, bold, italic, size, superscript in runs:
        if not text:
            continue
        run = paragraph.add_run(text)
        run.bold = True if bold_all else bold
        run.italic = italic
        if superscript:
            run.font.superscript = True
        if size and 6 <= size <= 48:
            run.font.size = Pt(round(size * 2) / 2)


def _add_table(document, rows: list[list[str]]):
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    table.autofit = True
    for row_index, row in enumerate(rows):
        for column_index in range(columns):
            value = row[column_index] if column_index < len(row) else ""
            cell = table.cell(row_index, column_index)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            run = paragraph.add_run(value)
            run.font.size = Pt(9.5)
            if row_index == 0 and value:
                run.bold = True
    document.add_paragraph()
    return table


def _add_image(document, block: dict, warnings: list[str]) -> None:
    data = block.get("image")
    if not data:
        return
    try:
        rect = fitz.Rect(block["bbox"])
        width_inches = max(1.0, min(6.0, rect.width / 72))
        document.add_picture(io.BytesIO(data), width=Inches(width_inches))
    except Exception:
        warnings.append("One embedded image could not be placed in the Word file.")


def _paragraphs_from_ocr(text: str, style: str = "Normal") -> list[tuple[str, str]]:
    paragraphs: list[tuple[str, str]] = []
    for chunk in re.split(r"\n\s*\n", text):
        lines = [line.strip() for line in chunk.splitlines() if line.strip()]
        if not lines:
            continue
        current: list[str] = []
        for line in lines:
            if BULLET_RE.match(line) or NUMBER_RE.match(line):
                if current:
                    paragraphs.append((" ".join(current), style))
                    current = []
                paragraphs.append((line, "li"))
            else:
                current.append(line)
        if current:
            paragraphs.append((" ".join(current), style))
    return paragraphs


def _add_text_block(document, block: dict, body_size: float, headers: set[str], warnings: list[str]) -> None:
    text, runs = _block_text(block)
    if not text or text in headers:
        return
    max_size = max((run[3] for run in runs), default=body_size)
    bold = bool(runs) and all(run[1] for run in runs if run[0].strip())
    kind, level = _classify(text, max_size, body_size, bold, len(block.get("lines", [])))

    if kind == "li":
        match = BULLET_RE.match(text) or NUMBER_RE.match(text)
        content = match.group(2) if match else text
        style = "List Number" if NUMBER_RE.match(text) else "List Bullet"
        paragraph = document.add_paragraph(style=style)
        markerless_runs = []
        removed = len(text) - len(content)
        consumed = 0
        for text_part, b, i, size, sup in runs:
            if consumed >= removed:
                markerless_runs.append((text_part, b, i, size, sup))
            consumed += len(text_part)
        _add_runs(paragraph, markerless_runs or [(content, False, False, body_size, False)])
        return

    if kind == "h" and level:
        paragraph = document.add_paragraph(style=HEADING_STYLE[level])
        _add_runs(paragraph, runs, bold_all=True)
        paragraph.paragraph_format.space_before = Pt(10 if level > 1 else 14)
        paragraph.paragraph_format.space_after = Pt(4)
        return

    paragraph = document.add_paragraph()
    _add_runs(paragraph, runs)
    paragraph.paragraph_format.space_after = Pt(6)
    if text.endswith((".", "!", "?", ":", ";")) and len(text) > 60:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def convert(ctx: ConversionContext) -> ConversionResult:
    return _convert(ctx, force_ocr=False)


def convert_ocr(ctx: ConversionContext) -> ConversionResult:
    return _convert(ctx, force_ocr=True)


def _convert(ctx: ConversionContext, *, force_ocr: bool) -> ConversionResult:
    doc = _open_pdf(ctx.input)
    warnings: list[str] = []
    if doc.page_count > MAX_PAGES:
        page_limit = MAX_PAGES
        warnings.append(f"Hanya {page_limit} halaman pertama yang diproses (batas keamanan).")
    else:
        page_limit = doc.page_count

    ocr_mode = "always" if force_ocr else str(ctx.options.get("ocr", "auto")).lower()
    language = str(ctx.options.get("language", "eng"))
    keep_tables = bool(ctx.options.get("tables", True)) and not force_ocr
    page_breaks = bool(ctx.options.get("page_breaks", False))
    body_size = _body_font_size(doc)
    headers = _header_footer_texts(doc)
    used_ocr_pages: list[int] = []
    two_column_pages: list[int] = []

    document = Document()
    section = document.sections[0]
    first_page = doc[0].rect
    section.page_width = Inches(first_page.width / 72)
    section.page_height = Inches(first_page.height / 72)
    section.left_margin = section.right_margin = Inches(0.8)
    section.top_margin = section.bottom_margin = Inches(0.8)
    document.core_properties.title = Path(ctx.original_name).stem

    try:
        for page_index in range(page_limit):
            page = doc[page_index]
            if page_index > 0 and page_breaks:
                document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

            page_text = page.get_text("text").strip()
            needs_ocr = force_ocr or ocr_mode == "always"
            if ocr_mode == "auto" and len(page_text) < 20:
                needs_ocr = True
            if ocr_mode == "never" and not force_ocr:
                needs_ocr = False

            if needs_ocr:
                if not ocr_available():
                    warnings.append(
                        f"Halaman {page_index + 1} tampaknya hasil pindai dan tidak memiliki lapisan teks. "
                        "OCR tidak tersedia, sehingga halaman ini dilewati."
                    )
                    continue
                try:
                    text = ocr_pdf_page(page, language=language)
                except Exception as exc:
                    warnings.append(f"OCR gagal pada halaman {page_index + 1}: {exc}")
                    continue
                if not text:
                    warnings.append(f"OCR tidak menemukan teks pada halaman {page_index + 1}.")
                    continue
                used_ocr_pages.append(page_index + 1)
                for content, kind in _paragraphs_from_ocr(text):
                    if kind == "li":
                        match = BULLET_RE.match(content) or NUMBER_RE.match(content)
                        if match:
                            style = "List Number" if NUMBER_RE.match(content) else "List Bullet"
                            document.add_paragraph(match.group(2), style=style)
                        else:
                            document.add_paragraph(content, style="List Bullet")
                    else:
                        document.add_paragraph(content)
                continue

            blocks = [
                block
                for block in page.get_text("dict").get("blocks", [])
                if block.get("type") in (0, 1)
            ]

            tables: list[tuple[tuple[float, float, float, float], list[list[str]]]] = []
            if keep_tables:
                try:
                    tables = _page_tables(ctx.input, page_index)
                except Exception:
                    warnings.append(f"Tabel pada halaman {page_index + 1} tidak dapat dianalisis.")

            text_blocks = [block for block in blocks if block.get("type") == 0]
            left = [b for b in text_blocks if fitz.Rect(b["bbox"]).x1 <= page.rect.width / 2 + 4]
            right = [b for b in text_blocks if fitz.Rect(b["bbox"]).x0 >= page.rect.width / 2 - 4]
            two_column = len(left) >= 2 and len(right) >= 2 and not tables
            if two_column:
                two_column_pages.append(page_index + 1)

            if two_column:
                ordered = sorted(left, key=lambda b: b["bbox"][1]) + sorted(right, key=lambda b: b["bbox"][1])
            else:
                ordered = sorted(blocks, key=lambda b: b["bbox"][1])

            table_rects = [fitz.Rect(bbox) for bbox, _ in tables]
            emitted: set[int] = set()

            if not two_column:
                interleaved: list[tuple[float, str, object]] = []
                for block in ordered:
                    interleaved.append((block["bbox"][1], block.get("type", 0), block))
                for index, (bbox, rows) in enumerate(tables):
                    interleaved.append((bbox[1], 2, (index, rows)))
                interleaved.sort(key=lambda item: item[0])
                ordered = [item[2] for item in interleaved]

            for item in ordered:
                if isinstance(item, tuple):
                    index, rows = item
                    if index in emitted:
                        continue
                    emitted.add(index)
                    _add_table(document, rows)
                    continue

                block = item
                if block.get("type") == 1:
                    _add_image(document, block, warnings)
                    continue

                block_rect = fitz.Rect(block["bbox"])
                if table_rects and any(
                    (block_rect & rect).get_area() > 0.5 * max(block_rect.get_area(), 1)
                    for rect in table_rects
                ):
                    continue
                _add_text_block(document, block, body_size, headers, warnings)
    finally:
        doc.close()

    if used_ocr_pages:
        warnings.insert(
            0,
            "Teks pada halaman "
            + ", ".join(str(page) for page in used_ocr_pages[:20])
            + " dihasilkan melalui OCR sehingga mungkin memiliki kesalahan pengenalan.",
        )
    if two_column_pages:
        warnings.append(
            "Tata letak dua kolom terdeteksi pada halaman "
            + ", ".join(str(page) for page in two_column_pages[:10])
            + ". Urutan baca disusun per kolom, tetapi posisi visual tidak sama persis."
        )
    if force_ocr:
        warnings.append("Mode OCR memaksa pengenalan teks pada semua halaman; tabel dan gambar tidak disalin.")
    elif keep_tables:
        warnings.append("Tabel disalin sebagai tabel Word biasa; penggabungan sel dan border kompleks dapat berubah.")

    if not document.paragraphs and not document.tables:
        raise ConversionError(
            "No content could be extracted from this PDF. If it is a scan, install an OCR engine and try again."
        )

    output = ctx.workdir / "converted.docx"
    document.save(output)
    filename = f"{Path(ctx.original_name).stem or 'document'}.docx"
    return ConversionResult(output, filename, DOCX_MIME, warnings)
