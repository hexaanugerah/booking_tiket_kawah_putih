"""PDF -> PPTX. Each PDF page becomes a slide with editable text boxes."""
from __future__ import annotations

from pathlib import Path

import fitz
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

from app.core.errors import ConversionError
from app.services.base import ConversionContext, ConversionResult
from app.services.pdf_ops import _open_pdf

PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
MAX_PAGES = 100
SLIDE_WIDTH_EMU = int(13.333 * 914400)
EMU_PER_POINT = 12700


def convert(ctx: ConversionContext) -> ConversionResult:
    doc = _open_pdf(ctx.input)
    include_images = bool(ctx.options.get("include_images", False))
    warnings: list[str] = []

    if doc.page_count > MAX_PAGES:
        warnings.append(f"Only the first {MAX_PAGES} pages were converted.")

    first_page = doc[0]
    page_width = first_page.rect.width or 595
    scale = (SLIDE_WIDTH_EMU / EMU_PER_POINT) / page_width

    prs = Presentation()
    prs.slide_width = Emu(SLIDE_WIDTH_EMU)
    prs.slide_height = Emu(int(doc[0].rect.height * scale * EMU_PER_POINT))

    try:
        for page_index in range(min(doc.page_count, MAX_PAGES)):
            page = doc[page_index]
            slide = prs.slides.add_slide(prs.slide_layouts[6])

            if include_images:
                pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                image_path = ctx.workdir / f"_slide-{page_index + 1}.png"
                pixmap.save(str(image_path))
                slide.shapes.add_picture(
                    str(image_path), 0, 0, width=prs.slide_width, height=prs.slide_height
                )

            blocks = [
                block
                for block in page.get_text("dict").get("blocks", [])
                if block.get("type") == 0
            ]
            if not blocks:
                warnings.append(f"Page {page_index + 1} has no extractable text layer.")
                continue

            for block in blocks:
                rect = fitz.Rect(block["bbox"])
                if rect.width < 2 or rect.height < 2:
                    continue
                left = Emu(int(rect.x0 * scale * EMU_PER_POINT))
                top = Emu(int(rect.y0 * scale * EMU_PER_POINT))
                width = Emu(int(max(rect.width, 20) * scale * EMU_PER_POINT))
                height = Emu(int(max(rect.height, 14) * scale * EMU_PER_POINT))
                box = slide.shapes.add_textbox(left, top, width, height)
                frame = box.text_frame
                frame.word_wrap = True
                frame.margin_left = 0
                frame.margin_right = 0
                frame.margin_top = 0
                frame.margin_bottom = 0

                paragraph = None
                for line in block.get("lines", []):
                    text = "".join(span.get("text", "") for span in line.get("spans", [])).strip()
                    if not text:
                        continue
                    paragraph = frame.add_paragraph() if paragraph is not None else frame.paragraphs[0]
                    run = paragraph.add_run()
                    run.text = text
                    size = max(span.get("size", 12) for span in line.get("spans", [{}]))
                    run.font.size = Pt(max(5, min(96, round(float(size) * scale, 1))))
                    bold = any(bool(span.get("flags", 0) & 16) for span in line.get("spans", []))
                    run.font.bold = bold
                    run.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
    finally:
        doc.close()

    output = ctx.workdir / "presentation.pptx"
    prs.save(output)
    if include_images:
        warnings.append("Rendered page images are placed behind the text boxes for layout fidelity.")
    else:
        warnings.append("Text is editable; complex shapes, charts and background graphics are not reproduced.")
    return ConversionResult(output, "presentation.pptx", PPTX_MIME, warnings)
