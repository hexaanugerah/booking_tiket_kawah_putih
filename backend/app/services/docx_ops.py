"""Word tools: DOCX -> PDF/TXT/HTML/Markdown, TXT/Markdown -> DOCX, merge, cleanup."""
from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.core.errors import ConversionError
from app.services.base import ConversionContext, ConversionResult
from app.services.capabilities import libreoffice_available
from app.services.html_ops import render_html_to_pdf, wrap_html
from app.services.html_to_docx import html_to_document
from app.services.office_ops import convert_with_soffice

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PDF_MIME = "application/pdf"


def _open_docx(path: Path) -> Document:
    try:
        return Document(str(path))
    except Exception as exc:
        raise ConversionError("The Word file could not be opened. It may be damaged or not a real DOCX.") from exc


def iter_block_items(document: Document):
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def _run_markdown(run) -> str:
    text = run.text
    if not text:
        return ""
    if run.bold and run.italic:
        return f"***{text}***"
    if run.bold:
        return f"**{text}**"
    if run.italic:
        return f"*{text}*"
    return text


def _table_markdown(table: Table) -> str:
    rows = []
    for row in table.rows:
        cells = [" ".join(cell.text.split()).replace("|", "\\|") for cell in row.cells]
        rows.append("| " + " | ".join(cells) + " |")
    if not rows:
        return ""
    columns = max(len(row.cells) for row in table.rows)
    separator = "| " + " | ".join(["---"] * columns) + " |"
    return "\n".join([rows[0], separator, *rows[1:]])


def docx_to_pdf(ctx: ConversionContext) -> ConversionResult:
    _open_docx(ctx.input)  # validates the file
    stem = Path(ctx.original_name).stem or "document"
    if libreoffice_available():
        output = convert_with_soffice(ctx.input, "pdf", ctx.workdir)
        return ConversionResult(output, f"{stem}.pdf", PDF_MIME, [])

    import mammoth

    try:
        with ctx.input.open("rb") as handle:
            result = mammoth.convert_to_html(handle)
    except Exception as exc:
        raise ConversionError("The Word document could not be converted.") from exc
    html = wrap_html(result.value, stem)
    output = render_html_to_pdf(html, ctx.workdir / "document.pdf")
    warnings = [
        "LibreOffice is not installed, so a simplified HTML-based renderer was used. "
        "Fonts, page breaks, columns and complex objects may differ from the original."
    ]
    if result.messages:
        warnings.append("Some Word features could not be reproduced by the fallback renderer.")
    return ConversionResult(output, f"{stem}.pdf", PDF_MIME, warnings)


def docx_to_txt(ctx: ConversionContext) -> ConversionResult:
    document = _open_docx(ctx.input)
    parts: list[str] = []
    for block in iter_block_items(document):
        if isinstance(block, Paragraph):
            parts.append(block.text)
        else:
            for row in block.rows:
                parts.append("\t".join(cell.text.replace("\n", " ") for cell in row.cells))
    text = "\n".join(parts).strip()
    if not text:
        raise ConversionError("The Word document contains no extractable text.")
    output = ctx.workdir / "document.txt"
    output.write_text(text, encoding="utf-8")
    return ConversionResult(output, f"{Path(ctx.original_name).stem or 'document'}.txt", "text/plain; charset=utf-8")


def docx_to_html(ctx: ConversionContext) -> ConversionResult:
    import mammoth

    try:
        with ctx.input.open("rb") as handle:
            result = mammoth.convert_to_html(handle)
    except Exception as exc:
        raise ConversionError("The Word document could not be converted to HTML.") from exc
    html = wrap_html(result.value, Path(ctx.original_name).stem)
    output = ctx.workdir / "document.html"
    output.write_text(html, encoding="utf-8")
    warnings = ["Some Word-specific formatting may be simplified in HTML."] if result.messages else []
    return ConversionResult(output, f"{Path(ctx.original_name).stem or 'document'}.html", "text/html; charset=utf-8", warnings)


def docx_to_markdown(ctx: ConversionContext) -> ConversionResult:
    document = _open_docx(ctx.input)
    lines: list[str] = []

    for block in iter_block_items(document):
        if isinstance(block, Table):
            table = _table_markdown(block)
            if table:
                lines.extend(["", table, ""])
            continue

        text = "".join(_run_markdown(run) for run in block.runs).strip() or block.text.strip()
        if not text:
            continue
        style = (block.style.name or "").lower()
        heading = re.match(r"heading (\d)", style)
        if heading:
            level = min(6, int(heading.group(1)))
            lines.extend([f"{'#' * level} {block.text.strip()}", ""])
        elif style.startswith("list bullet"):
            lines.append(f"- {block.text.strip()}")
        elif style.startswith("list number"):
            lines.append(f"1. {block.text.strip()}")
        elif style.startswith("quote") or style.startswith("intense quote"):
            lines.extend([f"> {block.text.strip()}", ""])
        else:
            lines.extend([text, ""])

    markdown = "\n".join(lines).strip()
    if not markdown:
        raise ConversionError("The Word document contains no convertible content.")
    output = ctx.workdir / "document.md"
    output.write_text(markdown + "\n", encoding="utf-8")
    return ConversionResult(output, f"{Path(ctx.original_name).stem or 'document'}.md", "text/markdown; charset=utf-8")


def txt_to_docx(ctx: ConversionContext) -> ConversionResult:
    text = ctx.input.read_text(encoding="utf-8", errors="replace")
    if not text.strip():
        raise ConversionError("The text file is empty.")
    document = Document()
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(11)
    for chunk in re.split(r"\n\s*\n", text):
        lines = chunk.splitlines()
        if not lines:
            continue
        paragraph = document.add_paragraph()
        for index, line in enumerate(lines):
            if index:
                paragraph.add_run().add_break()
            paragraph.add_run(line.strip())
    output = ctx.workdir / "document.docx"
    document.save(output)
    return ConversionResult(output, f"{Path(ctx.original_name).stem or 'document'}.docx", DOCX_MIME)


def markdown_to_docx(ctx: ConversionContext) -> ConversionResult:
    import markdown

    source = ctx.input.read_text(encoding="utf-8", errors="replace")
    if not source.strip():
        raise ConversionError("The Markdown file is empty.")
    html = markdown.markdown(source, extensions=["tables", "fenced_code", "sane_lists"])
    document = Document()
    document, warnings = html_to_document(html)
    output = ctx.workdir / "document.docx"
    document.save(output)
    return ConversionResult(output, f"{Path(ctx.original_name).stem or 'document'}.docx", DOCX_MIME, warnings)


def merge_docx(ctx: ConversionContext) -> ConversionResult:
    if len(ctx.inputs) < 2:
        raise ConversionError("Select at least two DOCX files to merge.")
    try:
        from docxcompose.composer import Composer

        master = _open_docx(ctx.inputs[0])
        composer = Composer(master)
        for path in ctx.inputs[1:]:
            composer.append(_open_docx(path))
        output = ctx.workdir / "merged.docx"
        composer.save(output)
        return ConversionResult(output, "merged.docx", DOCX_MIME)
    except ImportError:
        raise ConversionError("The docxcompose package is required to merge Word documents on this server.")


def clean_docx(ctx: ConversionContext) -> ConversionResult:
    document = _open_docx(ctx.input)
    font_name = str(ctx.options.get("font", "Calibri") or "Calibri")[:60]
    font_size = max(8, min(24, int(ctx.options.get("font_size", 11) or 11)))

    normal = document.styles["Normal"]
    normal.font.name = font_name
    normal.font.size = Pt(font_size)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    empty_streak = 0
    previous = None
    for paragraph in list(document.paragraphs):
        style_name = (paragraph.style.name or "").lower()
        is_heading = style_name.startswith("heading")
        if not paragraph.text.strip() and not paragraph._p.xpath(".//w:drawing"):
            empty_streak += 1
            if empty_streak >= 2:
                paragraph._element.getparent().remove(paragraph._element)
            continue
        empty_streak = 0
        if previous is not None:
            paragraph.paragraph_format.space_before = Pt(8 if is_heading else 0)
        if not is_heading and not style_name.startswith("list"):
            paragraph.paragraph_format.space_after = Pt(6)
        previous = paragraph

    output = ctx.workdir / "cleaned.docx"
    document.save(output)
    return ConversionResult(
        output,
        f"{Path(ctx.original_name).stem or 'document'}-clean.docx",
        DOCX_MIME,
        ["Formatting normalization adjusts paragraph spacing and body font; custom styles may change."],
    )
