"""Small HTML -> DOCX renderer used by the Markdown converter."""
from __future__ import annotations

from html.parser import HTMLParser

from docx import Document


class HtmlToDocx(HTMLParser):
    def __init__(self, document: Document):
        super().__init__(convert_charrefs=True)
        self.doc = document
        self.paragraph = None
        self.heading_level: int | None = None
        self.bold = 0
        self.italic = 0
        self.monospace = 0
        self.ignore_depth = 0
        self.list_stack: list[str] = []
        self.in_table = False
        self.table_rows: list[list[str]] = []
        self.current_row: list[str] | None = None
        self.cell_text: list[str] | None = None
        self.cell_bold = False
        self.first_row_bold = False
        self.warnings: list[str] = []

    # ------------------------------------------------------------------ helpers
    def _new_paragraph(self, style: str | None = None):
        self.paragraph = self.doc.add_paragraph(style=style) if style else self.doc.add_paragraph()

    def _flush(self) -> None:
        self.paragraph = None

    def _ensure_paragraph(self):
        if self.paragraph is None:
            self._new_paragraph()
        return self.paragraph

    def _build_table(self) -> None:
        rows = [row for row in self.table_rows if any(cell.strip() for cell in row)]
        self.in_table = False
        self.table_rows = []
        if not rows:
            return
        columns = max(len(row) for row in rows)
        table = self.doc.add_table(rows=len(rows), cols=columns)
        table.style = "Table Grid"
        for row_index, row in enumerate(rows):
            for column_index in range(columns):
                value = row[column_index] if column_index < len(row) else ""
                cell = table.cell(row_index, column_index)
                cell.text = ""
                run = cell.paragraphs[0].add_run(value)
                run.font.size = None
                if row_index == 0 and self.first_row_bold:
                    run.bold = True
        self.first_row_bold = False

    # ------------------------------------------------------------------ parser
    def handle_starttag(self, tag: str, attrs) -> None:
        if tag in ("script", "style", "head"):
            self.ignore_depth += 1
            return
        if tag.startswith("h") and tag[1:].isdigit() and 1 <= int(tag[1:]) <= 6:
            self._flush()
            self.heading_level = int(tag[1:])
            self._new_paragraph(f"Heading {self.heading_level}")
            return
        if tag == "p":
            self._flush()
            self._ensure_paragraph()
            return
        if tag in ("ul", "ol"):
            self._flush()
            self.list_stack.append(tag)
            return
        if tag == "li":
            self._flush()
            style = "List Number" if self.list_stack and self.list_stack[-1] == "ol" else "List Bullet"
            self._new_paragraph(style)
            return
        if tag == "table" and not self.in_table:
            self._flush()
            self.in_table = True
            self.table_rows = []
            return
        if tag == "tr" and self.in_table:
            self.current_row = []
            return
        if tag in ("td", "th") and self.in_table:
            self.cell_text = []
            self.cell_bold = tag == "th"
            if self.cell_bold and not self.table_rows and self.current_row == []:
                self.first_row_bold = True
            return
        if tag == "br":
            if self.paragraph is not None:
                self.paragraph.add_run().add_break()
            return
        if tag in ("strong", "b"):
            self.bold += 1
        elif tag in ("em", "i"):
            self.italic += 1
        elif tag == "code":
            self.monospace += 1
        elif tag == "pre":
            self._flush()
            self._new_paragraph()
            self.monospace += 1
        elif tag == "img":
            self.warnings.append("Images embedded in Markdown are not copied to the Word file.")
        elif tag == "hr":
            self._flush()

    def handle_endtag(self, tag: str) -> None:
        if tag in ("script", "style", "head"):
            self.ignore_depth = max(0, self.ignore_depth - 1)
            return
        if tag.startswith("h") and tag[1:].isdigit():
            self._flush()
            self.heading_level = None
        elif tag == "p":
            self._flush()
        elif tag == "li":
            self._flush()
        elif tag in ("ul", "ol"):
            if self.list_stack:
                self.list_stack.pop()
            self._flush()
        elif tag in ("td", "th") and self.in_table:
            if self.current_row is not None and self.cell_text is not None:
                self.current_row.append(" ".join("".join(self.cell_text).split()))
            self.cell_text = None
        elif tag == "tr" and self.in_table:
            if self.current_row:
                self.table_rows.append(self.current_row)
            self.current_row = None
        elif tag == "table" and self.in_table:
            self._build_table()
        elif tag in ("strong", "b"):
            self.bold = max(0, self.bold - 1)
        elif tag in ("em", "i"):
            self.italic = max(0, self.italic - 1)
        elif tag == "code":
            self.monospace = max(0, self.monospace - 1)
        elif tag == "pre":
            self.monospace = max(0, self.monospace - 1)
            self._flush()

    def handle_data(self, data: str) -> None:
        if self.ignore_depth or not data:
            return
        if self.in_table and self.cell_text is not None:
            self.cell_text.append(data)
            return
        if not data.strip() and self.paragraph is None:
            return
        paragraph = self._ensure_paragraph()
        run = paragraph.add_run(data)
        if self.heading_level:
            run.bold = True
        elif self.bold:
            run.bold = True
        if self.italic:
            run.italic = True
        if self.monospace:
            run.font.name = "Courier New"
            from docx.shared import Pt

            run.font.size = Pt(9.5)


def html_to_document(html: str, document: Document | None = None) -> tuple[Document, list[str]]:
    document = document or Document()
    parser = HtmlToDocx(document)
    parser.feed(html)
    parser.close()
    if parser.in_table:
        parser._build_table()
    return document, parser.warnings
